from __future__ import annotations

from openpyxl import load_workbook

from stig_audit_pro.core.result_model import CheckResult
from stig_audit_pro.reports.excel_writer import EXPECTED_SHEETS, write_excel_report


def test_excel_report_contains_expected_sheets_and_counts(tmp_path):
    results = [
        CheckResult(ip="192.0.2.1", hostname="sw1", vuln_id="V-1",
                    stig_family="IOSXE_L2", title="Finding", severity="high", status="Open"),
        CheckResult(ip="192.0.2.1", hostname="sw1", vuln_id="V-2",
                    stig_family="IOSXE_L2", title="Pass", severity="medium", status="NotAFinding"),
        CheckResult(ip="192.0.2.2", hostname="sw2", vuln_id="V-3",
                    stig_family="IOSXE_NDM", title="Manual", severity="low", status="Not_Reviewed"),
    ]
    destination = write_excel_report(results, tmp_path / "audit.xlsx")
    workbook = load_workbook(destination)
    assert tuple(workbook.sheetnames) == EXPECTED_SHEETS
    summary = {
        workbook["Executive Summary"].cell(row, 1).value:
        workbook["Executive Summary"].cell(row, 2).value
        for row in range(2, workbook["Executive Summary"].max_row + 1)
    }
    assert summary["Device count"] == 2
    assert summary["Total checks"] == 3
    assert summary["Open"] == 1
    assert workbook["All Findings"].max_row == 4
    assert workbook["Evidence Index"].freeze_panes == "A2"

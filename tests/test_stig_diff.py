from pathlib import Path

from openpyxl import load_workbook

from stig_audit_pro.stig.stig_diff import (
    RuleChange,
    compare_stig_releases,
    export_stig_diff_xlsx,
)
from stig_audit_pro.stig.xccdf_importer import parse_xccdf_file
from stig_audit_pro.stig.stig_metadata import StigBenchmarkMetadata, StigRuleMetadata


FIXTURES = Path(__file__).parent / "fixtures"


def _releases():
    old = parse_xccdf_file(FIXTURES / "stig_release_1.xml", "IOSXE_L2")
    new = parse_xccdf_file(FIXTURES / "stig_release_2.xml", "IOSXE_L2")
    return old, new


def test_normalized_diff_all_required_change_categories():
    old, new = _releases()
    diff = compare_stig_releases(old, new)
    by_vuln = {item.vuln_id: item for item in diff.rule_diffs}

    assert by_vuln["V-100"].classification == RuleChange.UNCHANGED
    assert by_vuln["V-200"].classification == RuleChange.REMOVED
    assert by_vuln["V-700"].classification == RuleChange.ADDED
    assert by_vuln["V-300"].changed_fields == ["severity"]
    assert by_vuln["V-400"].changed_fields == ["check_text"]
    assert by_vuln["V-500"].changed_fields == ["fix_text"]
    assert by_vuln["V-600"].changed_fields == ["rule_id"]
    assert by_vuln["V-600"].match_basis == "vuln_id"


def test_first_release_is_baseline_not_every_rule_changed():
    _old, new = _releases()
    diff = compare_stig_releases(None, new)
    assert diff.status.value == "NO_PREVIOUS_RELEASE"
    assert diff.rule_diffs == []
    assert diff.summary["added"] == 0


def test_procedure_normalization_ignores_line_ending_noise():
    old, new = _releases()
    new.rules[0].check_text = new.rules[0].check_text.replace("\n", "\r\n")
    new.rules[0].calculate_fingerprints()
    diff = compare_stig_releases(old, new)
    unchanged = next(item for item in diff.rule_diffs if item.vuln_id == "V-100")
    assert unchanged.classification == RuleChange.UNCHANGED


def test_stig_diff_excel_has_required_sheets(tmp_path):
    old, new = _releases()
    diff = compare_stig_releases(
        old, new, checks_dir=FIXTURES / "stig_diff_checks.yaml"
    )
    path = export_stig_diff_xlsx(diff, tmp_path / "diff.xlsx")
    workbook = load_workbook(path, read_only=True)
    assert workbook.sheetnames == [
        "Summary", "Added Rules", "Removed Rules", "Changed Rules",
        "YAML Impact", "All Rules",
    ]


def test_ambiguous_identifiers_are_not_silently_matched():
    old = StigBenchmarkMetadata(
        family="IOSXE_L2", benchmark_id="B", version="1",
        rules=[
            StigRuleMetadata(vuln_id="V-1", rule_id="OLD-A", stig_id="S-A"),
            StigRuleMetadata(vuln_id="V-1", rule_id="OLD-B", stig_id="S-B"),
        ],
    )
    new = StigBenchmarkMetadata(
        family="IOSXE_L2", benchmark_id="B", version="2",
        rules=[StigRuleMetadata(vuln_id="V-1", rule_id="NEW", stig_id="S-N")],
    )
    diff = compare_stig_releases(old, new)
    assert not any(item.match_basis == "vuln_id" for item in diff.rule_diffs)
    assert any(item.ambiguous for item in diff.rule_diffs)

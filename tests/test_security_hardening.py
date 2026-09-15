import logging
import zipfile

import pytest
from openpyxl import load_workbook

from stig_audit_pro.core.archive_safety import UnsafeArchiveError, validate_zip
from stig_audit_pro.core.result_model import CheckResult
from stig_audit_pro.logging_config import configure_logging
from stig_audit_pro.reports.audit_report import write_csv_report
from stig_audit_pro.reports.excel_writer import write_excel_report
from stig_audit_pro.reports.spreadsheet_safety import safe_spreadsheet_value
from stig_audit_pro.stig.cklb_writer import CklbError, load_cklb


def test_zip_slip_is_rejected(tmp_path):
    path = tmp_path / "bad.zip"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("../../outside.xml", "<Benchmark />")
    with zipfile.ZipFile(path) as archive:
        with pytest.raises(UnsafeArchiveError):
            validate_zip(archive)


@pytest.mark.parametrize("value", ["=1+1", "+cmd|' /C calc'!A0", "-1+2", "@SUM(A1:A2)", "  =HYPERLINK('x')"])
def test_spreadsheet_formula_values_are_neutralized(value):
    assert safe_spreadsheet_value(value).startswith("'")


def test_malformed_cklb_has_useful_error(tmp_path):
    path = tmp_path / "bad.cklb"
    path.write_text("{not json", encoding="utf-8")
    with pytest.raises(CklbError):
        load_cklb(path)


def test_report_writers_neutralize_formula_like_finding_text(tmp_path):
    result = CheckResult(ip="10.0.0.1", hostname="SW1", vuln_id="V-1", stig_family="L2", severity="cat2", status="Open", title="=HYPERLINK(\"bad\")")
    csv_path = write_csv_report([result], tmp_path / "findings.csv")
    assert "'=HYPERLINK" in csv_path.read_text(encoding="utf-8-sig")
    excel_path = write_excel_report([result], tmp_path / "findings.xlsx")
    workbook = load_workbook(excel_path, read_only=True)
    values = list(workbook["All Findings"].iter_rows(values_only=True))[1]
    assert values[9].startswith("'=")


def test_ssh_library_logging_cannot_emit_banners_at_info_level():
    configure_logging()
    assert logging.getLogger("paramiko.transport").getEffectiveLevel() >= logging.WARNING
    assert logging.getLogger("netmiko").getEffectiveLevel() >= logging.WARNING

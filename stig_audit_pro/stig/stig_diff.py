"""Normalized STIG release comparison and actionable YAML impact analysis."""

from __future__ import annotations

import csv
import json
from collections import Counter
from enum import Enum
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

import yaml
from pydantic import BaseModel, ConfigDict, Field, computed_field

from stig_audit_pro.core.models import CheckDefinition, CheckLibrary
from stig_audit_pro.reports.spreadsheet_safety import safe_spreadsheet_row
from stig_audit_pro.stig.stig_metadata import (
    RULE_FINGERPRINT_FIELDS,
    StigBenchmarkMetadata,
    StigRuleMetadata,
    normalize_identifier,
)


class RuleChange(str, Enum):
    ADDED = "ADDED"
    REMOVED = "REMOVED"
    UNCHANGED = "UNCHANGED"
    CHANGED = "CHANGED"


class YamlImpactType(str, Enum):
    NO_LOCAL_CHECK = "NO_LOCAL_CHECK"
    NEW_CHECK_REQUIRED = "NEW_CHECK_REQUIRED"
    METADATA_UPDATE = "METADATA_UPDATE"
    AUTOMATION_REVIEW_REQUIRED = "AUTOMATION_REVIEW_REQUIRED"
    FIX_GUIDANCE_CHANGED = "FIX_GUIDANCE_CHANGED"
    MANUAL_RULE_REVIEW = "MANUAL_RULE_REVIEW"
    RETIRE_CHECK_REVIEW = "RETIRE_CHECK_REVIEW"
    NO_ACTION_REQUIRED = "NO_ACTION_REQUIRED"
    AMBIGUOUS_MAPPING = "AMBIGUOUS_MAPPING"


class DiffStatus(str, Enum):
    COMPLETE = "COMPLETE"
    NO_PREVIOUS_RELEASE = "NO_PREVIOUS_RELEASE"


class DiffModel(BaseModel):
    model_config = ConfigDict(extra="forbid", use_enum_values=False)


class FieldDiff(DiffModel):
    field: str
    old_value: str = ""
    new_value: str = ""
    old_fingerprint: str = ""
    new_fingerprint: str = ""


class RuleDiff(DiffModel):
    classification: RuleChange
    family: str
    old_rule: StigRuleMetadata | None = None
    new_rule: StigRuleMetadata | None = None
    field_diffs: list[FieldDiff] = Field(default_factory=list)
    match_basis: str | None = None
    ambiguous: bool = False
    notes: list[str] = Field(default_factory=list)

    @computed_field
    @property
    def vuln_id(self) -> str:
        rule = self.new_rule or self.old_rule
        return rule.vuln_id if rule else ""

    @computed_field
    @property
    def old_rule_id(self) -> str:
        return self.old_rule.rule_id if self.old_rule else ""

    @computed_field
    @property
    def new_rule_id(self) -> str:
        return self.new_rule.rule_id if self.new_rule else ""

    @computed_field
    @property
    def changed_fields(self) -> list[str]:
        return [item.field for item in self.field_diffs]

    @property
    def change(self) -> RuleChange:
        return self.classification

    @property
    def change_type(self) -> RuleChange:
        return self.classification


class YamlImpact(DiffModel):
    family: str
    vuln_id: str
    old_rule_id: str = ""
    new_rule_id: str = ""
    change: RuleChange
    changed_fields: list[str] = Field(default_factory=list)
    yaml_file: str = ""
    check_id: str = ""
    check_type: str = ""
    automation_status: str = ""
    impact: YamlImpactType
    recommended_action: str
    old_release: str = ""
    new_release: str = ""
    current_rule_fingerprint: str = ""
    last_reviewed_stig_fingerprint: str = ""

    @property
    def classification(self) -> YamlImpactType:
        return self.impact


class StigDiff(DiffModel):
    status: DiffStatus = DiffStatus.COMPLETE
    family: str = ""
    benchmark_id: str = ""
    old_release: str = ""
    new_release: str = ""
    previous_rule_count: int = 0
    current_rule_count: int = 0
    rule_diffs: list[RuleDiff] = Field(default_factory=list)
    yaml_impacts: list[YamlImpact] = Field(default_factory=list)

    @property
    def rules(self) -> list[RuleDiff]:
        return self.rule_diffs

    @computed_field
    @property
    def summary(self) -> dict[str, int | str]:
        counts = Counter(item.classification.value for item in self.rule_diffs)
        impacts = Counter(item.impact.value for item in self.yaml_impacts)
        return {
            "status": self.status.value,
            "rules_previous": self.previous_rule_count,
            "rules_current": self.current_rule_count,
            "added": counts[RuleChange.ADDED.value],
            "removed": counts[RuleChange.REMOVED.value],
            "changed": counts[RuleChange.CHANGED.value],
            "unchanged": counts[RuleChange.UNCHANGED.value],
            "yaml_review_required": (
                impacts[YamlImpactType.AUTOMATION_REVIEW_REQUIRED.value]
                + impacts[YamlImpactType.MANUAL_RULE_REVIEW.value]
            ),
            "metadata_updates": impacts[YamlImpactType.METADATA_UPDATE.value],
            "new_checks_required": impacts[YamlImpactType.NEW_CHECK_REQUIRED.value],
            "retire_review": impacts[YamlImpactType.RETIRE_CHECK_REVIEW.value],
        }


class _MappedCheck:
    def __init__(self, check: CheckDefinition, path: str = "") -> None:
        self.check = check
        self.path = path


class StigDiffEngine:
    """Compare meaningful rule values; never perform a raw XML diff."""

    MATCH_FIELDS: tuple[str, ...] = ("vuln_id", "rule_id", "stig_id")

    def compare(
        self,
        previous: StigBenchmarkMetadata | None,
        current: StigBenchmarkMetadata,
        *,
        yaml_checks: Iterable[CheckDefinition | Mapping[str, Any]] | None = None,
        checks_dir: str | Path | None = None,
    ) -> StigDiff:
        current.calculate_fingerprints()
        if previous is None:
            diff = StigDiff(
                status=DiffStatus.NO_PREVIOUS_RELEASE,
                family=current.family,
                benchmark_id=current.benchmark_id,
                new_release=current.release_label,
                current_rule_count=len(current.rules),
            )
            # A baseline is deliberately not represented as N changed/added
            # rules. Coverage may still be calculated independently.
            return diff

        previous.calculate_fingerprints()
        pairs, old_only, new_only, ambiguity = self._match_rules(
            previous.rules, current.rules
        )
        changes: list[RuleDiff] = []
        for old_rule, new_rule, basis in pairs:
            field_diffs = self._field_diffs(old_rule, new_rule)
            changes.append(RuleDiff(
                classification=(RuleChange.CHANGED if field_diffs else RuleChange.UNCHANGED),
                family=current.family or previous.family,
                old_rule=old_rule,
                new_rule=new_rule,
                field_diffs=field_diffs,
                match_basis=basis,
            ))
        for old_rule in old_only:
            notes = ambiguity.get(("old", id(old_rule)), [])
            changes.append(RuleDiff(
                classification=RuleChange.REMOVED,
                family=previous.family,
                old_rule=old_rule,
                ambiguous=bool(notes),
                notes=notes,
            ))
        for new_rule in new_only:
            notes = ambiguity.get(("new", id(new_rule)), [])
            changes.append(RuleDiff(
                classification=RuleChange.ADDED,
                family=current.family,
                new_rule=new_rule,
                ambiguous=bool(notes),
                notes=notes,
            ))
        order = {
            RuleChange.ADDED: 0, RuleChange.REMOVED: 1,
            RuleChange.CHANGED: 2, RuleChange.UNCHANGED: 3,
        }
        changes.sort(key=lambda item: (
            order[item.classification], normalize_identifier(item.vuln_id),
            normalize_identifier(item.new_rule_id or item.old_rule_id),
        ))
        diff = StigDiff(
            family=current.family or previous.family,
            benchmark_id=current.benchmark_id or previous.benchmark_id,
            old_release=previous.release_label,
            new_release=current.release_label,
            previous_rule_count=len(previous.rules),
            current_rule_count=len(current.rules),
            rule_diffs=changes,
        )
        if yaml_checks is not None or checks_dir is not None:
            diff.yaml_impacts = analyze_yaml_impact(
                diff,
                mappings=yaml_checks,
                checks_dir=checks_dir,
            )
        return diff

    @classmethod
    def _match_rules(
        cls,
        old_rules: Sequence[StigRuleMetadata],
        new_rules: Sequence[StigRuleMetadata],
    ) -> tuple[
        list[tuple[StigRuleMetadata, StigRuleMetadata, str]],
        list[StigRuleMetadata],
        list[StigRuleMetadata],
        dict[tuple[str, int], list[str]],
    ]:
        unmatched_old = set(range(len(old_rules)))
        unmatched_new = set(range(len(new_rules)))
        pairs: list[tuple[StigRuleMetadata, StigRuleMetadata, str]] = []
        ambiguity: dict[tuple[str, int], list[str]] = {}
        for field_name in cls.MATCH_FIELDS:
            old_values: dict[str, list[int]] = {}
            new_values: dict[str, list[int]] = {}
            for index in unmatched_old:
                value = normalize_identifier(getattr(old_rules[index], field_name, ""))
                if value:
                    old_values.setdefault(value, []).append(index)
            for index in unmatched_new:
                value = normalize_identifier(getattr(new_rules[index], field_name, ""))
                if value:
                    new_values.setdefault(value, []).append(index)
            for value in sorted(set(old_values) & set(new_values)):
                old_matches, new_matches = old_values[value], new_values[value]
                if len(old_matches) != 1 or len(new_matches) != 1:
                    note = (
                        f"Ambiguous {field_name} match {value!r}: "
                        f"{len(old_matches)} previous and {len(new_matches)} current rules."
                    )
                    for index in old_matches:
                        ambiguity.setdefault(("old", id(old_rules[index])), []).append(note)
                    for index in new_matches:
                        ambiguity.setdefault(("new", id(new_rules[index])), []).append(note)
                    continue
                old_index, new_index = old_matches[0], new_matches[0]
                pairs.append((old_rules[old_index], new_rules[new_index], field_name))
                unmatched_old.remove(old_index)
                unmatched_new.remove(new_index)
        return (
            pairs,
            [old_rules[index] for index in sorted(unmatched_old)],
            [new_rules[index] for index in sorted(unmatched_new)],
            ambiguity,
        )

    @staticmethod
    def _field_diffs(
        old_rule: StigRuleMetadata, new_rule: StigRuleMetadata
    ) -> list[FieldDiff]:
        old_rule.calculate_fingerprints()
        new_rule.calculate_fingerprints()
        result: list[FieldDiff] = []
        for field_name in RULE_FINGERPRINT_FIELDS:
            old_hash = old_rule.field_fingerprints[field_name]
            new_hash = new_rule.field_fingerprints[field_name]
            if old_hash != new_hash:
                result.append(FieldDiff(
                    field=field_name,
                    old_value=str(getattr(old_rule, field_name, "") or ""),
                    new_value=str(getattr(new_rule, field_name, "") or ""),
                    old_fingerprint=old_hash,
                    new_fingerprint=new_hash,
                ))
        return result


def compare_stig_releases(
    previous: StigBenchmarkMetadata | None,
    current: StigBenchmarkMetadata,
    *,
    yaml_checks: Iterable[CheckDefinition | Mapping[str, Any]] | None = None,
    checks_dir: str | Path | None = None,
) -> StigDiff:
    return StigDiffEngine().compare(
        previous, current, yaml_checks=yaml_checks, checks_dir=checks_dir
    )


def _load_mappings(
    mappings: Iterable[CheckDefinition | Mapping[str, Any]] | None,
    checks_dir: str | Path | None,
) -> list[_MappedCheck]:
    loaded: list[_MappedCheck] = []
    for item in mappings or ():
        if isinstance(item, CheckDefinition):
            loaded.append(_MappedCheck(item, str(getattr(item, "check_file", ""))))
        else:
            payload = dict(item)
            path = str(payload.pop("check_file", payload.pop("yaml_file", "")))
            loaded.append(_MappedCheck(CheckDefinition.model_validate(payload), path))
    if checks_dir is not None:
        root = Path(checks_dir)
        paths = sorted(root.glob("*.yaml")) if root.is_dir() else [root]
        for path in paths:
            try:
                raw = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
                library = CheckLibrary.model_validate(raw)
            except (OSError, yaml.YAMLError, ValueError) as exc:
                raise ValueError(f"Invalid check library {path}: {exc}") from exc
            loaded.extend(_MappedCheck(check, path.as_posix()) for check in library.checks)
    return loaded


def _match_mappings(rule: RuleDiff, checks: list[_MappedCheck]) -> list[_MappedCheck]:
    candidates = [item for item in (rule.new_rule, rule.old_rule) if item is not None]
    for field_name in ("vuln_id", "rule_id", "stig_id"):
        values = {
            normalize_identifier(getattr(candidate, field_name, ""))
            for candidate in candidates
            if getattr(candidate, field_name, "")
        }
        matched = [
            item for item in checks
            if normalize_identifier(getattr(item.check, field_name, "")) in values
        ]
        if matched:
            return matched
    return []


def _automation_status(check: CheckDefinition) -> str:
    if check.automation_status:
        return check.automation_status
    return (
        "AUTOMATED"
        if check.automated and check.check_type != "manual_review"
        else "MANUAL_REVIEW"
    )


def _recommendation(
    impact: YamlImpactType, check: CheckDefinition | None, rule: RuleDiff
) -> str:
    check_type = check.check_type if check else "manual_review"
    messages = {
        YamlImpactType.NEW_CHECK_REQUIRED: (
            "Create a new manual_review starter check, then determine whether it can be automated."
        ),
        YamlImpactType.RETIRE_CHECK_REVIEW: (
            "This vulnerability is absent from the new release. Do not execute it by default "
            "for that release; retain the historical mapping for prior Audit Runs."
        ),
        YamlImpactType.METADATA_UPDATE: (
            "Update the local STIG metadata after confirming the identifier/title/severity change."
        ),
        YamlImpactType.AUTOMATION_REVIEW_REQUIRED: (
            f"Review the existing {check_type} automation against the new DISA check procedure. "
            "Do not consider the current YAML valid automatically."
        ),
        YamlImpactType.FIX_GUIDANCE_CHANGED: (
            "Review and display the new DISA fix guidance; audit automation is not changed automatically."
        ),
        YamlImpactType.MANUAL_RULE_REVIEW: (
            "Review the revised DISA procedure manually and update the manual-review guidance."
        ),
        YamlImpactType.NO_LOCAL_CHECK: (
            "Map this rule to a local check or document why it remains manual/out of scope."
        ),
        YamlImpactType.NO_ACTION_REQUIRED: (
            "No STIG-driven YAML change is required for this rule."
        ),
        YamlImpactType.AMBIGUOUS_MAPPING: (
            "Resolve the duplicate or conflicting identifiers explicitly; no mapping was guessed."
        ),
    }
    return messages[impact]


def analyze_yaml_impact(
    diff: StigDiff,
    checks_dir: str | Path | None = None,
    mappings: Iterable[CheckDefinition | Mapping[str, Any]] | None = None,
) -> list[YamlImpact]:
    checks = _load_mappings(mappings, checks_dir)
    impacts: list[YamlImpact] = []
    for rule in diff.rule_diffs:
        matched = _match_mappings(rule, checks)
        ambiguous = rule.ambiguous or len(matched) > 1
        if ambiguous:
            selected: list[_MappedCheck | None] = matched or [None]
            impact_type = YamlImpactType.AMBIGUOUS_MAPPING
        elif not matched:
            selected = [None]
            impact_type = (
                YamlImpactType.NEW_CHECK_REQUIRED
                if rule.classification == RuleChange.ADDED
                else YamlImpactType.NO_LOCAL_CHECK
            )
        else:
            selected = matched
            changed = set(rule.changed_fields)
            check = matched[0].check
            is_manual = not check.automated or check.check_type == "manual_review"
            if rule.classification == RuleChange.ADDED:
                impact_type = (
                    YamlImpactType.MANUAL_RULE_REVIEW
                    if is_manual else YamlImpactType.AUTOMATION_REVIEW_REQUIRED
                )
            elif rule.classification == RuleChange.REMOVED:
                impact_type = YamlImpactType.RETIRE_CHECK_REVIEW
            elif rule.classification == RuleChange.UNCHANGED:
                impact_type = YamlImpactType.NO_ACTION_REQUIRED
            elif "check_text" in changed:
                impact_type = (
                    YamlImpactType.MANUAL_RULE_REVIEW
                    if is_manual else YamlImpactType.AUTOMATION_REVIEW_REQUIRED
                )
            elif "fix_text" in changed:
                impact_type = YamlImpactType.FIX_GUIDANCE_CHANGED
            elif changed:
                impact_type = YamlImpactType.METADATA_UPDATE
            else:
                impact_type = YamlImpactType.NO_ACTION_REQUIRED
        for item in selected:
            check = item.check if item else None
            active_rule = rule.new_rule or rule.old_rule
            impacts.append(YamlImpact(
                family=rule.family,
                vuln_id=rule.vuln_id,
                old_rule_id=rule.old_rule_id,
                new_rule_id=rule.new_rule_id,
                change=rule.classification,
                changed_fields=rule.changed_fields,
                yaml_file=item.path if item else "",
                check_id=(check.check_id or check.vuln_id) if check else "",
                check_type=check.check_type if check else "",
                automation_status=_automation_status(check) if check else "NO_LOCAL_CHECK",
                impact=impact_type,
                recommended_action=_recommendation(impact_type, check, rule),
                old_release=diff.old_release,
                new_release=diff.new_release,
                current_rule_fingerprint=active_rule.rule_fingerprint if active_rule else "",
                last_reviewed_stig_fingerprint=(
                    check.last_reviewed_stig_fingerprint or "" if check else ""
                ),
            ))
    diff.yaml_impacts = impacts
    return impacts


def _jsonable(value: Any) -> Any:
    if isinstance(value, BaseModel):
        return value.model_dump(mode="json")
    if isinstance(value, Enum):
        return value.value
    return value


def export_stig_diff_json(diff: StigDiff, path: str | Path) -> Path:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    payload = {"schema_version": 1, **diff.model_dump(mode="json")}
    destination.write_text(
        json.dumps(payload, indent=2, sort_keys=True, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    return destination


_EXPORT_COLUMNS = (
    "Family", "Vuln ID", "Old Rule ID", "New Rule ID", "Change",
    "Changed Fields", "YAML File", "Check Type", "Automation Status",
    "Impact", "Recommended Action",
)


def _export_rows(diff: StigDiff) -> list[list[Any]]:
    by_key: dict[tuple[str, str, str], list[YamlImpact]] = {}
    for impact in diff.yaml_impacts:
        by_key.setdefault(
            (impact.vuln_id, impact.old_rule_id, impact.new_rule_id), []
        ).append(impact)
    rows: list[list[Any]] = []
    for rule in diff.rule_diffs:
        matches = by_key.get((rule.vuln_id, rule.old_rule_id, rule.new_rule_id), []) or [None]
        for impact in matches:
            rows.append([
                rule.family, rule.vuln_id, rule.old_rule_id, rule.new_rule_id,
                rule.classification.value, ", ".join(rule.changed_fields),
                impact.yaml_file if impact else "",
                impact.check_type if impact else "",
                impact.automation_status if impact else "",
                impact.impact.value if impact else "",
                impact.recommended_action if impact else "",
            ])
    return rows


def export_stig_diff_csv(diff: StigDiff, path: str | Path) -> Path:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(_EXPORT_COLUMNS)
        writer.writerows(safe_spreadsheet_row(row) for row in _export_rows(diff))
    return destination


def export_stig_diff_xlsx(diff: StigDiff, path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Metric", "Value"])
    summary.append(["Family", diff.family])
    summary.append(["Benchmark", diff.benchmark_id])
    summary.append(["Previous Release", diff.old_release])
    summary.append(["Current Release", diff.new_release])
    for key, value in diff.summary.items():
        summary.append([key.replace("_", " ").title(), value])

    rows = _export_rows(diff)
    sheets = {
        "Added Rules": RuleChange.ADDED.value,
        "Removed Rules": RuleChange.REMOVED.value,
        "Changed Rules": RuleChange.CHANGED.value,
        "YAML Impact": None,
        "All Rules": None,
    }
    for name, change in sheets.items():
        sheet = workbook.create_sheet(name)
        sheet.append(list(_EXPORT_COLUMNS))
        for row in rows:
            if change is None or row[4] == change:
                sheet.append(safe_spreadsheet_row(row))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column, width in enumerate((16, 16, 24, 24, 13, 30, 36, 26, 22, 30, 80), 1):
            sheet.column_dimensions[chr(64 + column)].width = width
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
    workbook.save(destination)
    return destination


# Common spelling used by callers.
export_stig_diff_excel = export_stig_diff_xlsx
ImpactClassification = YamlImpactType


__all__ = [
    "DiffStatus", "FieldDiff", "ImpactClassification", "RuleChange",
    "RuleDiff", "StigDiff", "StigDiffEngine", "YamlImpact",
    "YamlImpactType", "analyze_yaml_impact", "compare_stig_releases",
    "export_stig_diff_csv", "export_stig_diff_excel",
    "export_stig_diff_json", "export_stig_diff_xlsx",
]

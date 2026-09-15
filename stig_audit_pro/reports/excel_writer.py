"""Professional, filterable Excel audit reports."""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from stig_audit_pro.core.result_model import CheckResult
from stig_audit_pro.reports.spreadsheet_safety import safe_spreadsheet_row, safe_spreadsheet_value


EXPECTED_SHEETS = (
    "Executive Summary", "Device Summary", "All Findings", "CAT I", "CAT II",
    "CAT III", "Manual Review", "Errors", "Evidence Index",
)
FINDING_HEADERS = (
    "Run ID", "Device IP", "Hostname", "Vuln ID", "Rule ID", "STIG ID",
    "Family", "Severity", "Status", "Title", "Evaluation Reason",
    "Finding Details", "Comments", "Failed Objects", "Parser Warnings",
    "Commands Used", "Evidence Artifact IDs", "Evaluated At UTC",
)


def _get(value: Any, name: str, default: Any = "") -> Any:
    if value is None:
        return default
    if isinstance(value, dict):
        return value.get(name, default)
    return getattr(value, name, default)


def _severity_category(severity: str) -> str:
    value = severity.strip().lower().replace("_", " ")
    if value in {"high", "cat i", "cat1", "category i", "1"}:
        return "CAT I"
    if value in {"medium", "cat ii", "cat2", "category ii", "2"}:
        return "CAT II"
    if value in {"low", "cat iii", "cat3", "category iii", "3"}:
        return "CAT III"
    return ""


def _objects(items: Iterable[Any]) -> str:
    return "; ".join(
        f"{_get(item, 'object_type')}:{_get(item, 'object_name')}"
        + (f" ({_get(item, 'details')})" if _get(item, "details") else "")
        for item in items
    )


def _finding_row(result: CheckResult) -> list[Any]:
    timestamp = _get(result, "evaluated_at", None) or result.timestamp
    if isinstance(timestamp, datetime) and timestamp.tzinfo is not None:
        timestamp = timestamp.astimezone(timezone.utc).replace(tzinfo=None)
    return [
        _get(result, "run_id"), result.ip, result.hostname, result.vuln_id,
        _get(result, "rule_id"), _get(result, "stig_id"), result.stig_family,
        result.severity, result.status, result.title,
        _get(result, "evaluation_reason"), result.finding_details, result.comments,
        _objects(result.failed_objects), "; ".join(result.parser_warnings),
        "; ".join(result.commands_used),
        "; ".join(str(item) for item in _get(result, "evidence_artifact_ids", [])),
        timestamp,
    ]


def write_excel_report(
    results: Iterable[CheckResult],
    path: str | Path,
    *,
    audit_run: Any | None = None,
    evidence_artifacts: Iterable[Any] = (),
) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required for Excel reports") from exc

    result_list = list(results)
    artifact_list = list(evidence_artifacts)
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title in EXPECTED_SHEETS:
        workbook.create_sheet(title)

    navy, blue, pale = "17365D", "2F75B5", "D9EAF7"
    header_fill = PatternFill("solid", fgColor=navy)
    header_font = Font(color="FFFFFF", bold=True)

    def style_tabular(sheet, headers: Iterable[str], rows: Iterable[Iterable[Any]], table_name: str) -> None:
        header_list = list(headers)
        sheet.append(header_list)
        for row in rows:
            sheet.append(safe_spreadsheet_row(list(row)))
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        if sheet.max_row > 1:
            table = Table(displayName=table_name, ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showFirstColumn=False,
                showLastColumn=False, showRowStripes=True, showColumnStripes=False,
            )
            sheet.add_table(table)
        for column, header in enumerate(header_list, start=1):
            values = [str(sheet.cell(row, column).value or "") for row in range(1, min(sheet.max_row, 200) + 1)]
            width = min(60, max(10, len(str(header)) + 2, *(len(value) + 2 for value in values)))
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.sheet_view.showGridLines = False

    counts = Counter(result.status for result in result_list)
    severities = Counter(
        _severity_category(result.severity) for result in result_list
        if result.status == "Open"
    )
    reviewed = counts["NotAFinding"] + counts["Open"] + counts["Not_Applicable"]
    compliance = counts["NotAFinding"] / reviewed * 100.0 if reviewed else 0.0
    families = sorted({result.stig_family for result in result_list})
    summary_rows = [
        ("Run ID", _get(audit_run, "run_id")),
        ("Date (UTC)", _get(audit_run, "started_at") or datetime.now(timezone.utc).isoformat()),
        ("STIG families / versions", ", ".join(families)),
        ("Profile", _get(audit_run, "profile_name")),
        ("Device count", len({result.ip for result in result_list})),
        ("Total checks", len(result_list)), ("Open", counts["Open"]),
        ("NotAFinding", counts["NotAFinding"]),
        ("NotApplicable", counts["Not_Applicable"]),
        ("NotReviewed", counts["Not_Reviewed"]),
        ("Error", counts["Error"] + counts["Skipped"]),
        ("Compliance %", round(compliance, 2)),
        ("CAT I open", severities["CAT I"]),
        ("CAT II open", severities["CAT II"]),
        ("CAT III open", severities["CAT III"]),
    ]
    summary = workbook["Executive Summary"]
    summary["A1"] = "STIG Audit Pro — Executive Summary"
    summary["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor=navy)
    summary.merge_cells("A1:B1")
    for row in summary_rows:
        summary.append(safe_spreadsheet_row(list(row)))
    for row in range(2, summary.max_row + 1):
        summary.cell(row, 1).font = Font(bold=True, color=navy)
        if row % 2 == 0:
            for column in (1, 2):
                summary.cell(row, column).fill = PatternFill("solid", fgColor=pale)
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 48
    summary.freeze_panes = "A2"
    summary.sheet_view.showGridLines = False

    per_device: dict[tuple[str, str], Counter[str]] = defaultdict(Counter)
    for result in result_list:
        per_device[(result.ip, result.hostname)][result.status] += 1
    device_rows = [
        (ip, hostname, sum(statuses.values()), statuses["Open"],
         statuses["NotAFinding"], statuses["Not_Applicable"],
         statuses["Not_Reviewed"], statuses["Error"] + statuses["Skipped"])
        for (ip, hostname), statuses in sorted(per_device.items())
    ]
    style_tabular(
        workbook["Device Summary"],
        ("Device IP", "Hostname", "Total", "Open", "NotAFinding", "NotApplicable", "NotReviewed", "Errors"),
        device_rows, "DeviceSummaryTable",
    )
    style_tabular(workbook["All Findings"], FINDING_HEADERS, map(_finding_row, result_list), "AllFindingsTable")
    for category, table_name in (("CAT I", "CatITable"), ("CAT II", "CatIITable"), ("CAT III", "CatIIITable")):
        style_tabular(
            workbook[category], FINDING_HEADERS,
            (_finding_row(result) for result in result_list if _severity_category(result.severity) == category),
            table_name,
        )
    style_tabular(
        workbook["Manual Review"], FINDING_HEADERS,
        (_finding_row(result) for result in result_list if result.status == "Not_Reviewed"),
        "ManualReviewTable",
    )
    style_tabular(
        workbook["Errors"], FINDING_HEADERS,
        (_finding_row(result) for result in result_list if result.status in {"Error", "Skipped"}),
        "ErrorsTable",
    )
    evidence_headers = (
        "Artifact ID", "Device ID", "Command", "Relative Path", "SHA-256",
        "Collected At UTC", "Byte Length",
    )
    evidence_rows = [
        (_get(item, "artifact_id", _get(item, "id")), _get(item, "run_device_id"),
         _get(item, "command"), _get(item, "relative_path"), _get(item, "sha256"),
         _get(item, "collected_at"), _get(item, "byte_length"))
        for item in artifact_list
    ]
    style_tabular(workbook["Evidence Index"], evidence_headers, evidence_rows, "EvidenceIndexTable")
    for sheet in workbook.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.sheet_properties.tabColor = blue
    workbook.save(destination)
    return destination


__all__ = ["EXPECTED_SHEETS", "write_excel_report"]
